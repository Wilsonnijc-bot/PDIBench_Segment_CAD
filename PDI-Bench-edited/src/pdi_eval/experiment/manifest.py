#!/usr/bin/env python3
"""Match deformation workbook annotations to a verified source manifest."""

from __future__ import annotations

import argparse
import hashlib
import json
import re
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


ANNOTATION_SHEET = "Selected Annotations"
REQUIRED_COLUMNS = (
    "Dataset",
    "Video ID",
    "Original Instruction / Task",
    "Robot deformation (0/1)",
    "Forearm deformation",
    "Gripper deformation",
    "Upper arm deformation",
    "deformation degree(0-3)(1-minor,2 moderate, 3-severe)",
)


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def normalize_dataset(value: Any) -> str:
    normalized = " ".join(str(value or "").split()).lower()
    if normalized.startswith("lvp"):
        return "LVP_ROBOWM"
    if "2.5" in normalized:
        return "COSMOS2.5"
    if "cosmos3" in normalized or "cosmos 3" in normalized:
        return "COSMOS3"
    raise ValueError(f"unrecognized annotation dataset {value!r}")


def normalize_video_number(value: Any) -> str:
    match = re.search(r"(\d+)\s*$", str(value or ""))
    if match is None:
        raise ValueError(f"video ID has no trailing number: {value!r}")
    return f"{int(match.group(1)):04d}"


def optional_int(value: Any, *, field: str, row: int) -> int | None:
    if value is None or str(value).strip() == "":
        return None
    if isinstance(value, bool):
        result = int(value)
    elif isinstance(value, (int, float)) and float(value).is_integer():
        result = int(value)
    else:
        raise ValueError(f"row {row} field {field!r} is not an integer: {value!r}")
    return result


def load_annotations(path: Path) -> list[dict[str, Any]]:
    workbook = load_workbook(path, read_only=False, data_only=True)
    if ANNOTATION_SHEET not in workbook.sheetnames:
        raise ValueError(f"workbook is missing sheet {ANNOTATION_SHEET!r}")
    sheet = workbook[ANNOTATION_SHEET]
    headers = [cell.value for cell in sheet[2]]
    missing = [name for name in REQUIRED_COLUMNS if name not in headers]
    if missing:
        raise ValueError(f"annotation sheet is missing columns: {missing}")
    column = {name: index + 1 for index, name in enumerate(headers)}
    annotations = []
    for row in range(3, sheet.max_row + 1):
        video_id = sheet.cell(row, column["Video ID"]).value
        if video_id is None or str(video_id).strip() == "":
            continue
        raw_dataset = sheet.cell(row, column["Dataset"]).value
        dataset = normalize_dataset(raw_dataset)
        video_number = normalize_video_number(video_id)
        labels = {
            "robot_deformation": optional_int(
                sheet.cell(row, column["Robot deformation (0/1)"]).value,
                field="Robot deformation (0/1)",
                row=row,
            ),
            "upper_arm_deformation": optional_int(
                sheet.cell(row, column["Upper arm deformation"]).value,
                field="Upper arm deformation",
                row=row,
            ),
            "forearm_deformation": optional_int(
                sheet.cell(row, column["Forearm deformation"]).value,
                field="Forearm deformation",
                row=row,
            ),
            "gripper_deformation": optional_int(
                sheet.cell(row, column["Gripper deformation"]).value,
                field="Gripper deformation",
                row=row,
            ),
            "deformation_degree": optional_int(
                sheet.cell(
                    row,
                    column[
                        "deformation degree(0-3)(1-minor,2 moderate, 3-severe)"
                    ],
                ).value,
                field="deformation degree",
                row=row,
            ),
        }
        annotations.append(
            {
                "workbook_row": row,
                "raw_dataset": raw_dataset,
                "dataset": dataset,
                "video_id": str(video_id).strip(),
                "video_number": video_number,
                "task": sheet.cell(
                    row, column["Original Instruction / Task"]
                ).value,
                "labels": labels,
            }
        )
    keys = [(item["dataset"], item["video_number"]) for item in annotations]
    duplicates = sorted(key for key, count in Counter(keys).items() if count > 1)
    if duplicates:
        raise ValueError(f"duplicate dataset-qualified annotation rows: {duplicates}")
    return annotations


def source_index(manifest: dict[str, Any]) -> dict[tuple[str, str], dict[str, Any]]:
    grouped: dict[tuple[str, str], list[dict[str, Any]]] = defaultdict(list)
    for entry in manifest.get("videos", []):
        dataset = normalize_dataset(entry.get("dataset"))
        relative_path = entry.get("relative_path")
        if not relative_path:
            raise ValueError("source manifest entry is missing relative_path")
        number = normalize_video_number(Path(relative_path).stem)
        grouped[(dataset, number)].append(entry)
    ambiguous = sorted(key for key, entries in grouped.items() if len(entries) != 1)
    if ambiguous:
        raise ValueError(f"source manifest has ambiguous video keys: {ambiguous}")
    return {key: entries[0] for key, entries in grouped.items()}


def job_slug(dataset: str, video_number: str, sha256: str) -> str:
    return f"{dataset}_{video_number}-{sha256[:12]}"


def build_manifest(
    workbook_path: Path,
    source_manifest_path: Path,
    *,
    replays_per_dataset: int = 0,
    replay_all: bool = False,
) -> dict[str, Any]:
    if replays_per_dataset < 0:
        raise ValueError("replays_per_dataset must be non-negative")
    annotations = load_annotations(workbook_path)
    source_manifest = json.loads(source_manifest_path.read_text(encoding="utf-8"))
    sources = source_index(source_manifest)
    missing = sorted(
        (item["dataset"], item["video_number"])
        for item in annotations
        if (item["dataset"], item["video_number"]) not in sources
    )
    if missing:
        raise ValueError(f"annotated videos are missing from source manifest: {missing}")

    dataset_positions: dict[str, int] = defaultdict(int)
    records = []
    for annotation in annotations:
        key = (annotation["dataset"], annotation["video_number"])
        source = sources[key]
        dataset_position = dataset_positions[annotation["dataset"]]
        dataset_positions[annotation["dataset"]] += 1
        digest = str(source["sha256"])
        size = int(source["size_bytes"])
        staged_relative_path = f"{annotation['dataset']}/{annotation['video_number']}.mp4"
        records.append(
            {
                **annotation,
                "relative_path": f"{annotation['video_number']}.mp4",
                "staged_relative_path": staged_relative_path,
                "source_staged_relative_path": source["staged_relative_path"],
                "source_job_id": source.get("job_id"),
                "job_id": job_slug(
                    annotation["dataset"], annotation["video_number"], digest
                ),
                "sha256": digest,
                "size_bytes": size,
                "replay_selected": replay_all
                or dataset_position < replays_per_dataset,
            }
        )

    counts = Counter(record["dataset"] for record in records)
    payload = {
        "schema_version": 3,
        "pipeline": "v1-selection",
        "annotation": {
            "workbook_name": workbook_path.name,
            "workbook_sha256": sha256_file(workbook_path),
            "sheet": ANNOTATION_SHEET,
            "first_data_row": 3,
            "selected_row_count": len(records),
        },
        "source": {
            "manifest_name": source_manifest_path.name,
            "manifest_sha256": sha256_file(source_manifest_path),
            "video_count": source_manifest.get("video_count", len(sources)),
        },
        "video_count": len(records),
        "dataset_counts": dict(sorted(counts.items())),
        "total_size_bytes": sum(record["size_bytes"] for record in records),
        "replay_count": sum(record["replay_selected"] for record in records),
        "replays_per_dataset": replays_per_dataset,
        "replay_policy": "all" if replay_all else "first-n-per-dataset",
        "videos": records,
    }
    return payload


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", type=Path, required=True)
    parser.add_argument("--source-manifest", type=Path, required=True)
    parser.add_argument("--replays-per-dataset", type=int, default=0)
    parser.add_argument("--replay-all", action="store_true")
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args(argv)
    try:
        payload = build_manifest(
            args.workbook.resolve(),
            args.source_manifest.resolve(),
            replays_per_dataset=args.replays_per_dataset,
            replay_all=args.replay_all,
        )
    except (OSError, KeyError, TypeError, ValueError, json.JSONDecodeError) as exc:
        parser.error(str(exc))
    args.output.parent.mkdir(parents=True, exist_ok=True)
    temporary = args.output.with_suffix(".tmp.json")
    temporary.write_text(
        json.dumps(payload, indent=2, sort_keys=True) + "\n", encoding="utf-8"
    )
    temporary.replace(args.output)
    print(
        json.dumps(
            {
                "videos": payload["video_count"],
                "datasets": payload["dataset_counts"],
                "output": str(args.output.resolve()),
            },
            sort_keys=True,
        )
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
